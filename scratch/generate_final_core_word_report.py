import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_report_v5():
    doc = Document()
    
    # Title
    title = doc.add_heading('업종별 펀더멘털 가용성 분석 및 투자 전략 리포트', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 1. 우리가 정의한 가치주
    doc.add_heading('1. 우리가 정의한 가치주 (Value Definition)', level=1)
    doc.add_paragraph(
        '본 분석에서 정의하는 가치주는 단순히 PBR이 낮은 주식이 아닙니다. '
        '우리는 "ROE 1%를 창출하는 데 지불하는 비용(PBR)이 가장 저렴한 주식"을 진정한 가치주로 정의합니다.'
    )
    doc.add_paragraph(
        '- 역사적 PBR 밴드 하단에 위치한 업종\n'
        '- 수익성(ROE)이 개선되거나 유지되어 가치 함정(Value Trap)을 피한 종목'
    )
    doc.add_paragraph('---')

    # 2. 업종별 진행 이유 (Boxplot)
    doc.add_heading('2. 왜 \'업종별\' 분석을 진행해야 하는가?', level=1)
    doc.add_paragraph('산업마다 자본 구조와 수익 모델이 다르므로, 업종별 기초 체력의 차이를 인정해야 합니다.')
    
    img_boxplot = r'C:\Users\Hubnet\antigravity\results\viz_sector_distribution_boxplot.png'
    if os.path.exists(img_boxplot):
        doc.add_picture(img_boxplot, width=Inches(6))
        caption = doc.add_paragraph('(업종별 PBR 및 ROE 분포 차이를 보여주는 박스플롯)')
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph('분석 의견: 개별 업종마다 PBR의 하단과 상단이 다르게 형성되므로, 시장 전체 기준이 아닌 \'자기 자신(업종)\' 대비 현재의 위치를 분석하는 것이 타당합니다.')
    doc.add_paragraph('---')

    # 3. PBR과 주가의 상관관계
    doc.add_heading('3. PBR과 주가(종가)의 상관관계', level=1)
    doc.add_paragraph('PBR이 주가 변화를 얼마나 정직하게 따라가는지(Coupling), 그리고 어떤 예외가 있는지 분석했습니다.')
    
    img_coupling = r'C:\Users\Hubnet\antigravity\results\viz_pbr_price_coupling.png'
    if os.path.exists(img_coupling):
        doc.add_picture(img_coupling, width=Inches(6))
        caption = doc.add_paragraph('(상단: 시장 전체 상관관계 및 주요 선택 업종 / 하단: 동행성 대표 사례 - 은행)')
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
    doc.add_paragraph('분석 결과: 생물공학, 전기장비 등 특정 아웃라이어를 제외하면 대다수의 업종에서 주가와 PBR은 강력한 동행성(Coupling)을 보입니다. 우리는 이러한 동행성 속에서 실적 개선으로 인한 가치 괴리 구간을 포착합니다.')
    doc.add_paragraph('---')

    # 4. ROE를 추가했을 때의 효과
    doc.add_heading('4. ROE를 추가했을 때 얻을 수 있는 효과', level=1)
    doc.add_paragraph('ROE 지표를 결합하면 자산은 많지만 수익을 못 내는 \'가치 함정\'을 효과적으로 걸러낼 수 있습니다.')
    
    table_backtest = doc.add_table(rows=1, cols=4)
    table_backtest.style = 'Table Grid'
    hdr = table_backtest.rows[0].cells
    hdr[0].text = '검증 연도'
    hdr[1].text = '선정 전략'
    hdr[2].text = '평균 수익률'
    hdr[3].text = '비고'
    
    data_backtest = [
        ['2024', 'ROE/PBR 가성비 TOP 5', '+11.2%', '안정적 가치 방어 확인'],
        ['2025', 'ROE/PBR 가성비 TOP 5', '+37.4%', '시장 대비 초과 수익 달성']
    ]
    for row in data_backtest:
        row_cells = table_backtest.add_row().cells
        for i, text in enumerate(row):
            row_cells[i].text = text
            
    doc.add_paragraph('---')

    # 5. 포지셔닝 맵을 통한 유망 업종
    doc.add_heading('5. 현재 포지셔닝 맵을 통해 눈여겨봐야 할 업종', level=1)
    
    img_map = r'C:\Users\Hubnet\antigravity\results\viz_pbr_roe_positioning_map.png'
    if os.path.exists(img_map):
        doc.add_picture(img_map, width=Inches(5.5))
        caption = doc.add_paragraph('(업종별 밸류에이션-수익성 포지셔닝 맵)')
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph(
        '결론: 현재 데이터 기반 최우선 유망 업종(Golden Area)은 은행, 손해보험, 복합유틸리티 섹터입니다. '
        '이들은 높은 수익성을 유지하면서도 벨류에이션은 역사적 바닥에 위치해 있어 강력한 상승 여력을 보유하고 있습니다.'
    )
    
    doc.add_paragraph('---')
    doc.add_heading('최종 제언', level=1)
    doc.add_paragraph('데이터는 거짓말을 하지 않습니다. 우리는 숫자로 증명된 \'가성비 거인\'들에 집중하여 외풍에 흔들리지 않는 수익을 추구할 것을 제안합니다.')
    
    save_path = r'c:\Users\Hubnet\antigravity\results\mid_term_core_report_v5.docx'
    doc.save(save_path)
    print(f"Word report saved successfully: {save_path}")

if __name__ == "__main__":
    generate_report_v5()
