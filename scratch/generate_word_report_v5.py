from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os

def set_font(run, font_name, size=None, bold=False, color=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    if size:
        run.font.size = Pt(size)
    if bold:
        run.font.bold = True
    if color:
        run.font.color.rgb = color

def create_super_aggressive_report():
    doc = Document()
    
    # 기본 스타일 설정
    style = doc.styles['Normal']
    font = style.font
    font.name = '맑은 고딕'
    font.size = Pt(11)

    # === [제목 섹션] ===
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run('퀀트 투자 전략 결과 보고서 (초공격 모멘텀 턴어라운드)')
    set_font(run, '맑은 고딕', size=22, bold=True, color=RGBColor(192, 57, 43)) # 신뢰와 공격적인 레드 컬러
    
    info_para = doc.add_paragraph()
    info_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = info_para.add_run('작성일: 2026. 04. 13.\n분석 도구: Python / PostgreSQL / Gemini 3.1 Pro')
    set_font(run, '맑은 고딕', size=10, color=RGBColor(127, 140, 141))
    
    doc.add_paragraph('\n')

    # === [1. 개요] ===
    doc.add_heading('1. 개요 (Executive Summary)', level=1)
    doc.add_paragraph(
        '본 보고서는 시장의 급격한 변곡점과 주도주 교체 시그널을 극도로 민감하게 포착하기 위해 최적화된 [초공격형 모멘텀 턴어라운드] 모델의 개발 결과를 담고 있습니다. '
        '전통적인 가치 평가의 틀을 유지하되, 거래량의 폭발과 전문가 그룹의 집중적 관심(Report Burst)에 가중치를 대폭 상향하여 공격적인 수익을 지향합니다.'
    )

    # === [2. 강화된 모멘텀 전략 로직] ===
    doc.add_heading('2. 강화된 모멘텀 전략 로직 구성', level=1)
    
    p = doc.add_paragraph('2.1. 수급 모멘텀 및 리포트 Burst (Main Engine: 50%)', style='List Number')
    doc.add_paragraph(
        '• 거래량 폭발 (30%): 최근 1주간의 거래 점증을 포착하여 대규모 자금 유입의 실시간 신호를 잡아냅니다.',
        style='List Bullet 2'
    )
    doc.add_paragraph(
        '• 리포트 집중도 (20%): 최근 1개월 내 증권사 분석 결과가 집중된 종목을 포착하여 정보의 비대칭성 해소 시점을 공략합니다.',
        style='List Bullet 2'
    )
    
    p = doc.add_paragraph('2.2. 펀더멘털 안전판 (Safety Anchor: 40%)', style='List Number')
    doc.add_paragraph(
        '• 가치 지표 (20%): 업종 내 상대적 저평가를 확인하여 밸류에이션 버블 리스크를 필터링합니다.',
        style='List Bullet 2'
    )
    doc.add_paragraph(
        '• 수익성 지표 (20%): 실질적인 이익 창출 능력이 확인된 우량 기업만을 대상으로 하여 가치 함정을 회피합니다.',
        style='List Bullet 2'
    )

    p = doc.add_paragraph('2.3. 리스크 관리 자동화 (Advanced Filtering)', style='List Number')
    doc.add_paragraph(
        '• 투자의견 필터링 단축: 기존 3개월에서 1개월로 리서치 의견 필터링 기간을 단축하여, 최신의 시장 우려를 포트폴리오에 즉각 반영하도록 설계했습니다.',
        style='List Bullet 2'
    )

    # === [3. 전략 변경 사유 및 기대 효과] ===
    doc.add_heading('3. 전략 변경 사유 및 기대 효과', level=1)
    doc.add_paragraph(
        '조선 및 방산 섹터와 같이 거대한 업황 사이클이 시작될 때 나타나는 공통적인 징후(수급의 급증, 리포트의 동시 다발적 발행)를 놓치지 않기 위한 조치입니다.'
    )
    
    doc.add_paragraph(
        '• 기대 효과: 시장의 주도 테마가 형성되는 초기에 가장 빠르게 종목을 선점할 수 있으며, '
        '특히 하락장에서도 독자적인 시세를 내는 이른바 \'탈동조화(Decoupling)\' 주도주를 발굴하는 데 탁월한 성능을 보입니다.',
        style='List Bullet 2'
    )

    # === [4. 최종 확정 가중치 테이블] ===
    doc.add_heading('4. 최종 확정 가중치 (Super-Aggressive Configuration)', level=1)
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '평가 항목'
    hdr_cells[1].text = '가중치 (%)'
    hdr_cells[2].text = '전략적 역할'
    
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    data = [
        ('수급 (거래량 모멘텀)', '30%', '에너지 응축 및 상승 추진력'),
        ('리포트 폭발 (Burst)', '20%', '전문가/기관의 집중 관심도'),
        ('업종 내 저평가 (Value)', '20%', '하방 안전판 및 밸류 매력'),
        ('수익성 (ROE)', '20%', '실적 기반의 우량성 검증'),
        ('투자심리 (Sentiment)', '10%', '대중 심리 변곡점 보조')
    ]
    
    for item, weight, desc in data:
        row_cells = table.add_row().cells
        row_cells[0].text = item
        row_cells[1].text = weight
        row_cells[2].text = desc
        row_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph('\n')

    # === [5. 최종 결론] ===
    doc.add_heading('5. 최종 결론 및 실행 제언', level=1)
    doc.add_paragraph(
        '본 전략은 "바닥을 터치하고 상승 에너지가 분출되기 시작한" 종목을 가장 민첩하게 포착하는 시스템입니다. '
        '수급과 리포트 버스트에 기반한 초공격형 접근은 강한 상승 추세에서 압도적인 성과를 낼 것이며, '
        '철저한 리스크 필터링을 통해 시장의 일시적 소음으로부터 포트폴리오를 보호할 것입니다.'
    )

    # 파일 저장
    output_path = r'c:\Users\Hubnet\antigravity\results\strategy_report_super_aggressive.docx'
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Super-Aggressive Report generated at: {output_path}")

if __name__ == "__main__":
    create_super_aggressive_report()
