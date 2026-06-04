
import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_report():
    doc = Document()

    # Title
    title = doc.add_heading('퀀트 기반 업종 순환매 및 종목 발굴 시스템 구축 보고서', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1. 보고서의 목적
    doc.add_heading('1. 보고서의 목적', level=1)
    p1 = doc.add_paragraph(
        "본 시스템의 목적은 'MACD 히스토그램 국면 분석'을 통해 시장의 자금이 유입되는 주도 섹터를 선제적으로 포착하고, "
        "해당 섹터 내에서 반등 모멘텀이 가장 강력한 '대장주'를 과학적인 근거로 발굴하는 데 있습니다. "
        "단순한 가격 상승이 아닌, 에너지의 응축과 폭발(Z-Score & Golden Cross)을 포착하여 승률 높은 투자 기회를 제공합니다."
    )

    # 2. MACD 스윙 사이클에 따른 4대 국면 구조도
    doc.add_heading('2. MACD 스윙 사이클에 따른 4대 국면 구조도', level=1)
    doc.add_paragraph("MACD 히스토그램의 에너지 변화를 통해 현재 종목이 사이클의 어느 지점에 있는지를 진단합니다.")
    
    # Add Image
    img_path = r'C:\Users\Hubnet\.gemini\antigravity\brain\39c5bffe-b863-48c4-91ce-eefa5033e179\media__1777009248147.png'
    if os.path.exists(img_path):
        doc.add_picture(img_path, width=Inches(6))
    else:
        doc.add_paragraph("[이미지를 찾을 수 없습니다: " + img_path + "]")

    # Table for Phases
    table = doc.add_table(rows=5, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '국면 (Phase)'
    hdr_cells[1].text = '상태 정의 (MACD Histogram)'
    hdr_cells[2].text = '투자 심리 및 전략'

    phases = [
        ('1. 회복기', '지표 < 0, 상승 전환 (에너지 응축)', '최적의 저점 매수 타이밍 (MACD음수맥스)'),
        ('2. 상승기', '지표 > 0, 상승 추세 (에너지 폭발)', '강력한 추세 추종 및 불타기 구간 (골든크로스)'),
        ('3. 둔화기', '지표 > 0, 하락 전환 (에너지 분산)', '익절 및 비중 축소 고민 구간'),
        ('4. 하락기', '지표 < 0, 하락 추세 (에너지 소멸)', '관망 및 매도 구간 (공포 구간)')
    ]

    for p, d, s in phases:
        row_cells = table.add_row().cells
        row_cells[0].text = p
        row_cells[1].text = d
        row_cells[2].text = s

    # 3. 대시보드 핵심 점수 산출 예시
    doc.add_heading('3. 대시보드 핵심 점수 산출 예시 (Simulation)', level=1)
    
    doc.add_heading('[STEP 1] 중분류 활성 비율 (반도체 섹터 전체)', level=2)
    doc.add_paragraph("- 데이터: 전체 종목 113개 중 활성 종목(회복기+상승기) 67개")
    doc.add_paragraph("- 계산: (67 / 113) * 100 = 59.29%")
    doc.add_paragraph("- 결과: 중분류 점수 59.29점")

    doc.add_heading('[STEP 2] 소분류 가중 점수 (반도체 장비 업종)', level=2)
    doc.add_paragraph("- 데이터: 전체 종목 23개 중 회복기(P1) 9개, 상승기(P2) 7개")
    doc.add_paragraph("- 가중치 적용: 회복기 종목에 1.5배의 가중치 부여")
    doc.add_paragraph("  * 분자 계산: (9 * 1.5) + 7 = 20.5")
    doc.add_paragraph("- 계산: (20.5 / 23) * 100 = 89.13")
    doc.add_paragraph("- 결과: 소분류 점수 89.13점")

    doc.add_heading('[STEP 3] 최종 마스터 점수 (Master Score) 합산', level=2)
    doc.add_paragraph("- 소분류 반영분 (30%): 89.13 * 0.3 = 26.739")
    doc.add_paragraph("- 중분류 반영분 (70%): 59.29 * 0.7 = 41.503")
    doc.add_paragraph("- 최종 합산: 26.739 + 41.503 = 68.242")
    doc.add_paragraph("- 최종 결과: 68.24점 (마스터 점수)")

    # 4. 실전 추천 종목 발굴 전략
    doc.add_heading('4. 실전 추천 종목 발굴 전략 (Signal Engine)', level=1)
    
    doc.add_heading('[전략 1] MACD 음수맥스 (바닥 탈출형)', level=2)
    doc.add_paragraph("- 대상: 오늘 막 회복기(Phase 1)에 진입한 종목")
    doc.add_paragraph("- Z-Score 에너지 필터: 최근 20거래일 MACD 히스토그램 변동 폭 분석, Z-Score <= -1.5인 종목 추출")
    doc.add_paragraph("  * 의미: 에너지가 이례적으로 깊은 과매도 구간(심해)까지 털린 종목의 강력한 기술적 반등 포착")

    doc.add_heading('[전략 2] 진성 골든크로스 (추세 추종형)', level=2)
    doc.add_paragraph("- 대상: 오늘 막 상승기(Phase 2)에 진입한 종목")
    doc.add_paragraph("- 진성 신호 필터: 반드시 이전 국면이 회복기(1)였던 종목만 인정")
    doc.add_paragraph("  * 의미: 바닥권(음수)에서 힘을 비축해 0선을 뚫고 올라오는 진짜 추세 전환 종목 포착")

    # 5. 대시보드 정렬 및 관리 원칙
    doc.add_heading('5. 대시보드 정렬 및 관리 원칙', level=1)
    doc.add_paragraph("1. 섹터 테이블: 마스터 점수 -> 회복기 비중 순으로 정렬하여 주도주 선별")
    doc.add_paragraph("2. 종목 리스트: 현재 국면(회복기 우선) -> 진입일(최근 순) -> 시가총액 순으로 정렬")
    doc.add_paragraph("3. 종목 정화: 스팩(SPAC) 종목 자동 제외")

    # Footer
    doc.add_paragraph("\n\n2026-04-24\n퀀트 분석 시스템 구축팀", style='Normal').alignment = WD_ALIGN_PARAGRAPH.RIGHT

    save_path = r'C:\Users\Hubnet\antigravity\Sector_Rotation_Analysis_Report.docx'
    doc.save(save_path)
    print(f"Report saved to: {save_path}")

if __name__ == "__main__":
    create_report()
