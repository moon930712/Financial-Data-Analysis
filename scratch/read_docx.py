import docx

doc_path = r"C:\Users\Hubnet\antigravity\순환매_분석_보고서.docx"
doc = docx.Document(doc_path)

out = []
for p in doc.paragraphs:
    out.append(p.text)

for table in doc.tables:
    out.append("\n--- Table ---")
    for row in table.rows:
        row_text = []
        for cell in row.cells:
            row_text.append(cell.text.replace("\n", " "))
        out.append(" | ".join(row_text))
    out.append("--- End Table ---\n")

with open(r"C:\Users\Hubnet\antigravity\scratch\docx_text.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(out))
