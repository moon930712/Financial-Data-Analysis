import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re
import os

md_path = r"C:\Users\Hubnet\.gemini\antigravity\brain\448cd88b-00aa-4259-a8e1-b69d680b77e8\sector_rotation_report.md"
out_path = r"C:\Users\Hubnet\antigravity\순환매_분석_보고서.docx"

doc = docx.Document()

# Styles
style = doc.styles['Normal']
font = style.font
font.name = 'Malgun Gothic'
font.size = Pt(11)

with open(md_path, 'r', encoding='utf-8') as f:
    lines = f.readlines()

def add_paragraph_with_bold(text, style='Normal'):
    p = doc.add_paragraph(style=style)
    # Split by ** for bold
    parts = re.split(r'\*\*(.*?)\*\*', text)
    for i, part in enumerate(parts):
        run = p.add_run(part)
        if i % 2 == 1: # Odd indices are inside **
            run.bold = True
    return p

in_table = False
table_data = []

for line in lines:
    line = line.strip()
    if not line:
        continue
        
    if line == '---':
        continue
        
    if line.startswith('|'):
        in_table = True
        if '---' in line:
            continue # skip separator
        row_data = [cell.strip() for cell in line.split('|')[1:-1]]
        table_data.append(row_data)
        continue
    else:
        if in_table:
            # Create table
            table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
            table.style = 'Table Grid'
            for r_idx, row in enumerate(table_data):
                for c_idx, cell_text in enumerate(row):
                    cell = table.cell(r_idx, c_idx)
                    # Handle bold in table cells
                    parts = re.split(r'\*\*(.*?)\*\*', cell_text)
                    p = cell.paragraphs[0]
                    for i, part in enumerate(parts):
                        run = p.add_run(part)
                        if i % 2 == 1:
                            run.bold = True
            in_table = False
            table_data = []
            
    if line.startswith('# '):
        p = doc.add_heading(level=0)
        run = p.add_run(line[2:])
        run.font.name = 'Malgun Gothic'
    elif line.startswith('## '):
        p = doc.add_heading(level=1)
        run = p.add_run(line[3:])
        run.font.name = 'Malgun Gothic'
    elif line.startswith('### '):
        p = doc.add_heading(level=2)
        run = p.add_run(line[4:])
        run.font.name = 'Malgun Gothic'
    elif line.startswith('> '):
        p = add_paragraph_with_bold(line[2:], style='Intense Quote')
    elif line.startswith('* '):
        p = add_paragraph_with_bold(line[2:], style='List Bullet')
    elif re.match(r'^\d+\.\s', line):
        text = re.sub(r'^\d+\.\s', '', line)
        p = add_paragraph_with_bold(text, style='List Number')
    else:
        add_paragraph_with_bold(line)

if in_table:
    # Create table if file ends with table
    table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
    table.style = 'Table Grid'
    for r_idx, row in enumerate(table_data):
        for c_idx, cell_text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            parts = re.split(r'\*\*(.*?)\*\*', cell_text)
            p = cell.paragraphs[0]
            for i, part in enumerate(parts):
                run = p.add_run(part)
                if i % 2 == 1:
                    run.bold = True

doc.save(out_path)
print(f"Word document saved to: {out_path}")
