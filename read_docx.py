import docx
doc = docx.Document('Sector_Rotation_Midterm_Report.docx')
with open('docx_content.txt', 'w', encoding='utf-8') as f:
    for p in doc.paragraphs:
        f.write(p.text + '\n')
    f.write('---TABLES---\n')
    for table in doc.tables:
        for row in table.rows:
            f.write('\t'.join([cell.text for cell in row.cells]) + '\n')
