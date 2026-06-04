import sys
import subprocess

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    import seaborn as sns
except ImportError:
    install('seaborn')
    import seaborn as sns

try:
    import matplotlib.pyplot as plt
except ImportError:
    install('matplotlib')
    import matplotlib.pyplot as plt

import pandas as pd
import numpy as np
import docx
from docx.shared import Inches
import os
import warnings

warnings.filterwarnings('ignore')

# Font setting for Korean on Windows
plt.rc('font', family='Malgun Gothic')
plt.rcParams['axes.unicode_minus'] = False

# 1. Chart 1: Holding Period Comparison (5 vs 8 vs 14)
days = ['5일 보유', '8일 보유', '14일 보유']
median_returns = [1.40, 2.76, 4.09]
win_rates = [57.3, 62.0, 63.2]

fig, ax1 = plt.subplots(figsize=(8, 5))
color = 'tab:red'
ax1.set_xlabel('보유 기간', fontsize=12)
ax1.set_ylabel('중앙값 수익률 (%)', color=color, fontsize=12)
bars = ax1.bar(days, median_returns, color=color, alpha=0.6, width=0.5)
ax1.tick_params(axis='y', labelcolor=color)
ax1.set_ylim(0, 5)

for bar in bars:
    yval = bar.get_height()
    ax1.text(bar.get_x() + bar.get_width()/2, yval + 0.1, f'{yval}%', ha='center', va='bottom', fontweight='bold')

ax2 = ax1.twinx()
color = 'tab:blue'
ax2.set_ylabel('승률 (%)', color=color, fontsize=12)
line = ax2.plot(days, win_rates, color=color, marker='o', linewidth=2, markersize=8)
ax2.tick_params(axis='y', labelcolor=color)
ax2.set_ylim(50, 70)

for i, txt in enumerate(win_rates):
    ax2.annotate(f'{txt}%', (days[i], win_rates[i] + 0.5), ha='center', color=color, fontweight='bold')

plt.title('보유 기간별 2국면 종목 성과 우상향 차트', fontsize=15, fontweight='bold')
fig.tight_layout()
plt.savefig('chart1_holding.png', dpi=300)
plt.close()

# 2. Chart 2: Jump Theme Phase Returns
phases = ['1국면', '2국면', '3국면', '4국면']
jump_medians = [2.62, 10.13, 2.28, 6.14]
jump_wins = [57.5, 76.8, 57.8, 67.3]

fig, ax1 = plt.subplots(figsize=(8, 5))
color = 'tab:red'
ax1.set_ylabel('중앙값 수익률 (%)', color=color, fontsize=12)
bars = ax1.bar(phases, jump_medians, color=['gray', 'tab:red', 'gray', 'gray'], alpha=0.7, width=0.5)
ax1.tick_params(axis='y', labelcolor=color)
ax1.set_ylim(0, 12)

for bar in bars:
    yval = bar.get_height()
    ax1.text(bar.get_x() + bar.get_width()/2, yval + 0.2, f'{yval}%', ha='center', va='bottom', fontweight='bold')

ax2 = ax1.twinx()
color = 'tab:blue'
ax2.set_ylabel('승률 (%)', color=color, fontsize=12)
line = ax2.plot(phases, jump_wins, color=color, marker='o', linewidth=2, markersize=8)
ax2.tick_params(axis='y', labelcolor=color)
ax2.set_ylim(50, 85)

for i, txt in enumerate(jump_wins):
    ax2.annotate(f'{txt}%', (phases[i], jump_wins[i] + 1), ha='center', color=color, fontweight='bold')

plt.title('급등 테마(61~100위 -> 1~10위) 국면별 14일 성과', fontsize=15, fontweight='bold')
fig.tight_layout()
plt.savefig('chart2_phase.png', dpi=300)
plt.close()

# 3. Chart 3: Heatmap Transition
data = np.array([
    [1.92],
    [1.87],
    [0.78],
    [1.85],
    [0.55]
])
ylabels = ['최상위(1~10)', '상위(11~30)', '중위(31~60)', '중하위(61~100)', '최하위(101~150)']
xlabels = ['오늘 최상위 진입 시 수익률']

plt.figure(figsize=(6, 4))
ax = sns.heatmap(data, annot=True, fmt=".2f", cmap="Reds", cbar=False, 
                 xticklabels=xlabels, yticklabels=ylabels, annot_kws={"size": 14, "weight": "bold"})
for t in ax.texts: t.set_text(t.get_text() + " %")
plt.title('어제 순위별 -> 최상위 진입 시 중앙값 수익률', fontsize=12, fontweight='bold')
plt.yticks(rotation=0)
plt.tight_layout()
plt.savefig('chart3_heatmap.png', dpi=300)
plt.close()

print("Charts generated.")

# 4. Update Word Document
try:
    doc = docx.Document('Sector_Rotation_Midterm_Report_v3.docx')

    for i, p in enumerate(doc.paragraphs):
        if '3.2. 보유 기간별' in p.text:
            r = p.insert_paragraph_before('').add_run()
            r.add_picture('chart3_heatmap.png', width=Inches(4.0))
            break

    for i, p in enumerate(doc.paragraphs):
        if p.text.startswith('4. 중하위'):
            r = p.insert_paragraph_before('').add_run()
            r.add_picture('chart1_holding.png', width=Inches(5.5))
            break
            
    for i, p in enumerate(doc.paragraphs):
        if p.text.startswith('5. 결론 및 향후 과제'):
            r = p.insert_paragraph_before('').add_run()
            r.add_picture('chart2_phase.png', width=Inches(5.5))
            break

    doc.save('Sector_Rotation_Midterm_Report_v4.docx')
    print("Word document V4 saved successfully.")
except Exception as e:
    import traceback
    traceback.print_exc()
